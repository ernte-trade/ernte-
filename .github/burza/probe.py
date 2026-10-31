import urllib.request, re, http.cookiejar, datetime, io, subprocess, sys
subprocess.run([sys.executable,"-m","pip","install","-q","--user","--break-system-packages","openpyxl"],check=False);import site,importlib;sys.path.append(site.getusersitepackages());importlib.invalidate_caches()
import openpyxl
cj=http.cookiejar.CookieJar();op=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
UA={"User-Agent":"Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126 Safari/537.36","Accept-Language":"en-GB,en;q=0.9"}
op.open(urllib.request.Request("https://ahdb.org.uk/cereals-oilseeds/futures-prices",headers=UA),timeout=30).read()
e=datetime.date.today();a=e-datetime.timedelta(days=100)
for ck in ("94","71"):
    u=f"https://ahdb.org.uk/FuturesCustomReportExport/API?id=84399&currency=17&timescale=Daily&startDate={a}&endDate={e}&commodityKey={ck}&convertToTonnes=true"
    try:
        r=op.open(urllib.request.Request(u,headers=dict(UA,Referer="https://ahdb.org.uk/cereals-oilseeds/futures-prices")),timeout=60);b=r.read()
        print("DL",ck,r.status,r.headers.get("Content-Type"),r.headers.get("Content-Disposition"),len(b),b[:60])
        try:
            wb=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True)
            for ws in wb.worksheets[:3]:
                print("SHEET",ws.title)
                for i,row in enumerate(ws.iter_rows(values_only=True)):
                    if i<14 or i>ws.max_row-3: print("  ",row[:14])
        except Exception as x: print("xlsx?",x, b[:400])
    except Exception as x: print("DL",ck,x)
